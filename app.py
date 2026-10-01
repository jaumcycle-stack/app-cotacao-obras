import streamlit as st
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import io

st.set_page_config(page_title="Automação de Compras", page_icon="🏗️")

st.title("🏗️ Gerador de Mapa de Cotação")
st.write("Faça o upload do modelo Excel e dos orçamentos para gerar o mapa automatizado.")

# Interface de Upload
modelo_excel = st.file_uploader("1. Anexe o Modelo Excel (MPC - Modelo.xlsx)", type=["xlsx"])
fornecedores_pdfs = st.file_uploader("2. Anexe os PDFs dos Fornecedores (Máx 3)", type=["pdf"], accept_multiple_files=True)

if st.button("Gerar Mapa Final"):
    if not modelo_excel or len(fornecedores_pdfs) == 0:
        st.warning("Por favor, anexe o modelo Excel e pelo menos um orçamento em PDF.")
    else:
        try:
            # Carregar o Excel na memória
            wb = openpyxl.load_workbook(modelo_excel)
            ws = wb.active
            
            # --- LÓGICA DE PREENCHIMENTO (Baseada no seu modelo) ---
            # Aqui entrará a biblioteca de leitura de PDF (ex: pdfplumber) para extrair os valores reais.
            # Para esta demonstração, aplicamos as formatações e regras exigidas.
            
            # Alterar Cabeçalho
            ws.cell(row=2, column=5, value="MATERIAL EXTRAÍDO DO PDF")
            
            # Identificação do Vencedor (Exemplo visual prático)
            green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            font_winner = Font(name="Segoe UI", size=11, bold=True, color="006100")
            
            # Simulação de marcação do vencedor (Fornecedor 2)
            ws.cell(row=8, column=9).fill = green_fill
            ws.cell(row=8, column=9).font = font_winner
            ws.cell(row=8, column=10).fill = green_fill
            ws.cell(row=8, column=10).font = font_winner
            
            # --- CRIAÇÃO DA ÁREA DE ANEXOS ---
            start_row = 24
            ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row, end_column=12)
            header_cell = ws.cell(row=start_row, column=2)
            header_cell.value = "ÁREA DE ANEXOS - ORÇAMENTOS ORIGINAIS"
            header_cell.fill = PatternFill(start_color="002060", fill_type="solid")
            header_cell.font = Font(color="FFFFFF", bold=True)
            header_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            ws.merge_cells(start_row=start_row+1, start_column=2, end_row=start_row+1, end_column=12)
            ws.cell(row=start_row+1, column=2, value="Para anexar: Inserir -> Texto/Objeto -> Criar do Arquivo -> Marcar 'Exibir como ícone'").alignment = Alignment(horizontal="center")
            
            # Preparar ficheiro para download
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            st.success("Mapa de Cotação gerado com sucesso!")
            st.download_button(
                label="📥 Baixar Mapa Preenchido",
                data=output,
                file_name="Mapa_Cotacao_Pronto.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            st.error(f"Erro ao processar o ficheiro: {e}")
